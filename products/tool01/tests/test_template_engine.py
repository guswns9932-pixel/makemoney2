import shutil

import pytest
from pptx import Presentation

from src import template_engine as te


def test_extract_placeholders_from_sample(sample_pptx):
    placeholders = te.extract_placeholders(sample_pptx)
    assert set(placeholders) == {"수료번호", "이름", "과정명", "수료일", "기관명"}


def test_extract_placeholders_missing_file(tmp_path):
    with pytest.raises(te.TemplateError):
        te.extract_placeholders(tmp_path / "no_such.pptx")


def test_extract_placeholders_wrong_extension(tmp_path):
    p = tmp_path / "not_a_template.txt"
    p.write_text("{{x}}")
    with pytest.raises(te.TemplateError):
        te.extract_placeholders(p)


# TEST 09: 원본 보호 - render_pptx가 템플릿 원본을 수정하지 않는지 확인
def test_render_pptx_does_not_modify_original(sample_pptx, tmp_path):
    original_mtime = sample_pptx.stat().st_mtime
    original_size = sample_pptx.stat().st_size

    out = tmp_path / "rendered.pptx"
    mapping = {
        "수료번호": "CERT-2026-001",
        "이름": "홍길동",
        "과정명": "AI 업무자동화 실무과정",
        "수료일": "2026-08-20",
        "기관명": "ABC 교육원",
    }
    te.render_pptx(sample_pptx, out, mapping)

    assert sample_pptx.stat().st_mtime == original_mtime
    assert sample_pptx.stat().st_size == original_size
    assert out.exists()


# TEST 10: 한글 치환 후 깨짐 없이 반영되는지
def test_render_pptx_replaces_placeholders_correctly(sample_pptx, tmp_path):
    out = tmp_path / "rendered.pptx"
    mapping = {
        "수료번호": "CERT-2026-099",
        "이름": "김민지",
        "과정명": "데이터분석 기초과정",
        "수료일": "2026-08-19",
        "기관명": "한빛 직무교육센터",
    }
    te.render_pptx(sample_pptx, out, mapping)

    prs = Presentation(str(out))
    all_text = []
    for slide in prs.slides:
        for shape in slide.shapes:
            if shape.has_text_frame:
                all_text.append(shape.text_frame.text)
    joined = "\n".join(all_text)

    for value in mapping.values():
        assert value in joined
    assert "{{" not in joined  # Placeholder가 모두 치환되었는지


# Placeholder가 아닌 문구는 변경되지 않아야 한다
def test_render_pptx_preserves_non_placeholder_text(sample_pptx, tmp_path):
    out = tmp_path / "rendered.pptx"
    mapping = {"이름": "테스트", "수료번호": "X", "과정명": "X", "수료일": "X", "기관명": "X"}
    te.render_pptx(sample_pptx, out, mapping)

    prs = Presentation(str(out))
    all_text = "\n".join(
        s.text_frame.text for slide in prs.slides for s in slide.shapes if s.has_text_frame
    )
    assert "Certificate of Completion" in all_text
    assert "위 사람은 소정의 교육과정을 성실히 이수하였으므로 이 증서를 수여합니다." in all_text
